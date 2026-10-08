from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parent

def create_flow():
    im=Image.new('RGB',(1600,1850),'white');d=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',44)
    def text(cx,cy,label):
        lines=label.split('\n');height=len(lines)*52
        for i,line in enumerate(lines):
            box=d.textbbox((0,0),line,font=font)
            d.text((cx-(box[2]-box[0])/2,cy-height/2+i*52),line,font=font,fill='black')
    def rect(box,label):
        d.rectangle(box,outline='black',width=3)
        text((box[0]+box[2])/2,(box[1]+box[3])/2,label)
    def diamond(cx,cy,w,h,label):
        d.polygon([(cx,cy-h/2),(cx+w/2,cy),(cx,cy+h/2),(cx-w/2,cy)],outline='black',width=3)
        text(cx,cy,label)
    def arrow(points,label=None,pos=None):
        d.line(points,fill='black',width=3)
        x,y=points[-1];px,py=points[-2]
        if y>py:d.polygon([(x,y),(x-10,y-18),(x+10,y-18)],fill='black')
        elif y<py:d.polygon([(x,y),(x-10,y+18),(x+10,y+18)],fill='black')
        elif x>px:d.polygon([(x,y),(x-18,y-10),(x-18,y+10)],fill='black')
        else:d.polygon([(x,y),(x+18,y-10),(x+18,y+10)],fill='black')
        if label:
            box=d.textbbox(pos,label,font=font);d.rectangle((box[0]-3,box[1]-3,box[2]+3,box[3]+3),fill='white');d.text(pos,label,font=font,fill='black')
    rect((410,10,1190,100),'Click button Checkout')
    arrow([(800,100),(800,140)])
    rect((340,140,1260,255),'Lưu checkout → POST /api/orders\nXác thực và kiểm tra tạo trùng')
    arrow([(800,255),(800,300)])
    diamond(800,410,650,220,'Dữ liệu\nhợp lệ?')
    arrow([(1125,410),(1210,410)],'Không',(1125,285))
    rect((1210,350,1590,490),'Báo lỗi\nSửa checkout')
    arrow([(800,520),(800,565)],'Có',(825,520))
    diamond(800,670,670,210,'Có Customer\ntheo email?')
    arrow([(465,670),(275,670),(275,795)],'Có',(280,600))
    arrow([(1135,670),(1325,670),(1325,795)],'Chưa có',(1240,600))
    rect((10,795,540,920),'Lấy customer_id\nhiện có')
    rect((1060,795,1590,920),'Tạo Customer\ntừ checkout')
    arrow([(275,920),(275,975),(650,975),(650,1020)])
    arrow([(1325,920),(1325,975),(950,975),(950,1020)])
    rect((260,1020,1340,1145),'Đọc giá variant và phí giao realtime\nsub_total = tổng price × quantity')
    arrow([(800,1145),(800,1190)])
    rect((300,1190,1300,1300),'Kiểm tra shipping và COD\norder_total = sub_total + shipping_fee')
    arrow([(800,1300),(800,1340)])
    diamond(800,1440,650,200,'Hàng và giá\nhợp lệ?')
    arrow([(475,1440),(370,1440)],'Không',(385,1360))
    rect((10,1370,370,1510),'Báo lỗi hoặc\nxác nhận lại giá')
    arrow([(800,1540),(800,1580)],'Có',(825,1540))
    rect((260,1580,1340,1690),'Ghi Order và Items trong transaction\nTrừ kho → hoàn tất checkout → commit')
    arrow([(800,1690),(800,1730)])
    rect((260,1730,1340,1840),'Chuyển Order data\nGET đơn đã lưu và hiển thị thông tin')
    out=ROOT/'order-plan-assets'/'create-order-flow.png';out.parent.mkdir(exist_ok=True);im.save(out)
    return out
# Các bảng hiện có chỉ liệt kê các cột liên quan đến Create Order.
schemas=[
 ('users','Hiện có','Xác thực người thao tác và quyền xem đơn.',[
 ('id','bigint','PK'),('email','varchar(254)','Email đăng nhập'),('status','varchar(20)','Trạng thái tài khoản')]),
 ('shops','Hiện có','Shop bán hàng và xử lý từng đơn.',[
 ('id','bigint','PK'),('owner_user_id','bigint','FK users.id'),('name','varchar(100)','Tên shop'),('status','varchar(20)','Trạng thái bán')]),
 ('products','Hiện có','Kiểm tra sản phẩm thuộc shop và còn bán.',[
 ('id','bigint','PK'),('shop_id','bigint','FK shops.id'),('owner_id','bigint','FK users.id'),('title','varchar(180)','Tên sản phẩm'),('status','varchar(30)','Trạng thái bán')]),
 ('product_variants','Hiện có','Nguồn giá và tồn kho realtime.',[
 ('id','bigint','PK'),('product_id','bigint','FK products.id'),('sku','varchar(64)','UNIQUE'),('variant_key','varchar(255)','Khóa biến thể'),('price','decimal(12,2)','Giá hiện tại'),('stock_quantity','int unsigned','Tồn kho'),('status','varchar(20)','Trạng thái biến thể')]),
 ('CheckOutToken','Hiện có','Dữ liệu checkout khách đã chốt.',[
 ('id','bigint','PK'),('user_id','bigint','FK users.id'),('checkout_token','uuid','UNIQUE'),('items','json','product_id, variant_id, quantity'),('version','int unsigned','Kiểm soát chỉnh sửa'),('is_completed','boolean','Đã tạo đơn hay chưa'),('shipping_address','text','Địa chỉ nháp'),('shipping_method','text','Rate đã chọn mỗi shop'),('payment_method','json','COD đã chọn'),('total','text','Báo giá nháp để so sánh')]),
 ('payment_methods','Hiện có','Cấu hình COD của chủ shop.',[
 ('id','bigint','PK'),('user_id','bigint','FK users.id'),('name','varchar(100)','Tên phương thức'),('payment_data','json','type, description, instructions'),('is_active','boolean','Đang hoạt động'),('is_deleted','boolean','Đã xóa mềm')]),
 ('shipping_methods','Hiện có','Phương thức giao hàng của shop.',[
 ('id','bigint','PK'),('shop_id','bigint','FK shops.id'),('name','varchar(100)','Tên phương thức'),('code','varchar(50)','Mã phương thức'),('status','varchar(20)','Phải active')]),
 ('shipping_rates','Hiện có','Phí và số ngày giao của phương thức.',[
 ('id','bigint','PK'),('shipping_method_id','bigint','FK shipping_methods.id'),('fixed_fee','decimal(12,2)','Phí mỗi order shop'),('min_delivery_days','int unsigned','Số ngày tối thiểu'),('max_delivery_days','int unsigned','Số ngày tối đa')]),
 ('countries','Hiện có','Quốc gia shop hỗ trợ vận chuyển.',[
 ('id','bigint','PK'),('shop_id','bigint','FK shops.id'),('name','varchar(100)','Tên quốc gia'),('country_code','char(2)','Mã ISO; model hiện tại varchar(2)'),('phone_code','varchar(8)','Mã điện thoại')]),
 ('shipping_rate_countries','Hiện có','Nối rate với quốc gia giao được.',[
 ('shipping_rate_id','bigint','PK ghép và FK shipping_rates.id'),('country_id','bigint','PK ghép và FK countries.id')]),
 ('geo_countries','Tạo mới','Danh mục quốc gia chung cho địa chỉ.',[
 ('id','bigint','PK'),('code','char(2)','UNIQUE'),('name','varchar(100)','Tên quốc gia'),('phone_code','varchar(8)','Mã điện thoại')]),
 ('provinces','Tạo mới','Danh mục tỉnh thuộc quốc gia chung.',[
 ('id','bigint','PK'),('country_id','bigint','FK geo_countries.id'),('code','varchar(20)','UNIQUE(country_id, code)'),('name','varchar(100)','Tên tỉnh hoặc bang')]),
 ('customers','Tạo mới','Khách hàng riêng với User; tìm theo email.',[
 ('id','bigint','PK'),('email','varchar(254)','Email checkout'),('email_key','varchar(254)','UNIQUE; trim và lowercase'),('first_name','varchar(100)','Tên khách'),('last_name','varchar(100)','Họ khách'),('phone','varchar(30)','Điện thoại'),('country_id','bigint','FK geo_countries.id'),('province_id','bigint','FK provinces.id; nullable'),('city','varchar(100)','Thành phố'),('ward','varchar(100)','Phường hoặc xã; nullable'),('street','varchar(200)','Đường hoặc địa chỉ chi tiết'),('house_number','varchar(50)','Số nhà; nullable theo quốc gia'),('apartment','varchar(100)','Căn hộ; nullable'),('postal_code','varchar(20)','Mã bưu chính; theo quốc gia')]),
 ('order_groups','Tạo mới','Gom đơn nhiều shop và chống tạo trùng.',[
 ('id','bigint','PK'),('checkout_id','bigint','FK CheckOutToken.id; UNIQUE'),('user_id','bigint','FK users.id; quyền xem đơn'),('customer_id','bigint','FK customers.id'),('request_id','uuid','UNIQUE(user_id, request_id)'),('currency','char(3)','VND')]),
 ('orders','Tạo mới','Một đơn của một shop.',[
 ('id','bigint','PK'),('order_group_id','bigint','FK order_groups.id'),('customer_id','bigint','FK customers.id'),('shop_id','bigint','FK shops.id'),('order_code','varchar(50)','UNIQUE'),('order_name','varchar(100)','Tên đơn hiển thị'),('status','varchar(20)','Khởi tạo pending'),('sub_total','decimal(18,2)','Tổng line_total'),('shipping_fee','decimal(18,2)','Phí giao hàng'),('order_total','decimal(18,2)','sub_total + shipping_fee'),('currency','char(3)','VND'),('shipping_method_id','bigint','FK shipping_methods.id'),('shipping_rate_id','bigint','FK shipping_rates.id'),('shipping_method_name','varchar(100)','Tên tại lúc đặt'),('shipping_method_code','varchar(50)','Mã tại lúc đặt'),('min_delivery_days','int unsigned','Lưu từ rate'),('max_delivery_days','int unsigned','Lưu từ rate'),('estimated_delivery_from','date','Ngày dự kiến sớm nhất'),('estimated_delivery_to','date','Ngày dự kiến muộn nhất')]),
 ('order_items','Tạo mới','Các dòng hàng và giá tại lúc đặt.',[
 ('id','bigint','PK'),('order_id','bigint','FK orders.id'),('product_id','bigint','FK products.id'),('product_variant_id','bigint','FK product_variants.id'),('product_name','varchar(180)','Tên tại lúc đặt'),('sku','varchar(64)','SKU tại lúc đặt'),('variant_data','json','Size, màu tại lúc đặt'),('image_url','text','Ảnh; nullable'),('quantity','int unsigned','Số lượng checkout'),('unit_price','decimal(18,2)','Giá realtime lúc đặt'),('line_total','decimal(18,2)','unit_price × quantity')]),
 ('order_addresses','Tạo mới','Địa chỉ giao của đơn được lưu tách cột.',[
 ('id','bigint','PK'),('order_id','bigint','FK orders.id; UNIQUE'),('recipient_first_name','varchar(100)','Tên người nhận'),('recipient_last_name','varchar(100)','Họ người nhận'),('email','varchar(254)','Email liên hệ lúc đặt'),('phone','varchar(30)','Điện thoại lúc đặt'),('country_id','bigint','FK geo_countries.id'),('country_code','char(2)','Mã lúc đặt'),('country_name','varchar(100)','Tên lúc đặt'),('province_id','bigint','FK provinces.id; nullable'),('province_code','varchar(20)','Mã lúc đặt; nullable'),('province_name','varchar(100)','Tên lúc đặt; nullable'),('city','varchar(100)','Thành phố'),('ward','varchar(100)','Phường hoặc xã; nullable'),('street','varchar(200)','Đường hoặc địa chỉ chi tiết'),('house_number','varchar(50)','Số nhà; nullable theo quốc gia'),('apartment','varchar(100)','Căn hộ; nullable'),('postal_code','varchar(20)','Mã bưu chính; theo quốc gia')]),
 ('order_payments','Tạo mới','Phương thức và trạng thái thu tiền của đơn.',[
 ('id','bigint','PK'),('order_id','bigint','FK orders.id; UNIQUE'),('payment_method_id','bigint','FK payment_methods.id'),('method_type','varchar(30)','cod'),('method_name','varchar(100)','Tên tại lúc đặt'),('description','text','Mô tả; nullable'),('instructions','text','Hướng dẫn; nullable'),('amount','decimal(18,2)','Khớp order_total'),('currency','char(3)','VND'),('status','varchar(20)','Khởi tạo unpaid'),('paid_at','datetime','Nullable khi chưa thu tiền')])
]

doc=Document();s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.65)
for st in doc.styles:
    if st.type==1 or st.type==2:
        st.font.name='Arial';st.font.size=Pt(14);st.font.color.rgb=RGBColor(0,0,0)
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    doc.styles[name].font.size=Pt(22)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing=1.0
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.space_after=Pt(7)
def p(text): return doc.add_paragraph(text)
def h(text): doc.add_heading(text,1)
def tab(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,x in zip(t.rows[0].cells,headers):c.text=x
    for row in rows:
        for c,x in zip(t.add_row().cells,row):c.text=x
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');m.append(e)
            tc.append(m)
            if i==0:
                e=OxmlElement('w:shd');e.set(qn('w:fill'),'E8E8E8');tc.append(e)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1
                for r in para.runs:r.font.size=Pt(14);r.font.color.rgb=RGBColor(0,0,0);r.bold=i==0
    p('')

p('Kế hoạch Create Order').style='Title'
h('I Luồng API Create Order')
fp=p('');fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=fp.add_run();run.add_picture(str(create_flow()),width=Inches(7.2))
run._r.xpath('.//wp:docPr')[0].set('descr','Lưu đồ Create Order với bước xử lý hình chữ nhật và quyết định hình thoi, nối bằng mũi tên.')
doc.add_page_break()
h('II Workplan')
def step(title,items):
    doc.add_heading(title,2)
    for item in items:doc.add_paragraph(item,'List Bullet')
step('1 Tạo migration',[
'geo_countries, provinces, customers.',
'order_groups, orders, order_addresses, order_items, order_payments.',
'Thêm khóa ngoại, unique và index; seed quốc gia và tỉnh; chạy migration trên database test.'])
step('2 Tạo model',[
'Tạo model tương ứng với 8 bảng mới.',
'Khai báo belongsTo, hasMany và hasOne trong models/index.js.'])
step('3 Tạo service',[
'Customer service: tìm bằng email_key; chưa có thì tạo Customer.',
'Order calculation service: đọc giá variant, quantity checkout, phí vận chuyển và COD; tính lại tổng.',
'Order service: tạo group, order, items, address và payment trong transaction; trừ kho và hoàn tất checkout.',
'Order query service: lấy thông tin đơn và kiểm tra quyền xem.'])
step('4 Tạo validator controller và router',[
'Validator: checkout_token, version, request_id; kiểm tra địa chỉ khi xác nhận.',
'Controller: gọi service, trả kết quả và mã lỗi.',
'Router: khai báo các API sau và gắn middleware xác thực.'])
p('Các API cần tạo:')
doc.add_paragraph('POST /api/orders — tạo đơn từ checkout.','List Bullet')
doc.add_paragraph('GET /api/orders/:id — lấy chi tiết một đơn.','List Bullet')
doc.add_paragraph('GET /api/order-groups/:id — lấy toàn bộ đơn trong một lần mua.','List Bullet')
p('POST nhận checkout_token, version, request_id. Giá và phí do backend đọc lại; retry cùng request_id trả đơn đã tạo.')
step('5 Bổ sung dữ liệu checkout',[
'Thêm street, house_number, apartment và ward vào form, normalizer, validator và autosave.',
'Ánh xạ country/province; giữ địa chỉ tách cột trong Customer và Order address.'])
step('6 Thêm button Checkout',[
'Đợi autosave thành công rồi gọi POST /api/orders.',
'Hiện trạng thái chờ, chặn double click và giữ request_id khi retry.',
'Giá hoặc phí thay đổi: hiện báo giá mới để khách xác nhận lại.'])
step('7 Tạo màn Order data',[
'Tạo route /orders/:id và /order-groups/:id; gọi API GET tương ứng.',
'Hiển thị order_name, địa chỉ, shipping method, khoảng ngày giao, COD, items và tổng tiền.',
'Mỗi shop có một order; COD khởi tạo unpaid.'])
step('8 Kiểm thử',[
'Customer có/chưa có; một/nhiều shop; tổng tiền đúng.',
'Hết hàng, giá/phí đổi, COD tắt; rollback khi có lỗi.',
'Double click, mất mạng sau commit, quyền xem đơn và giao diện sáng/tối.'])
doc.add_page_break()
h('III Bảng dữ liệu')
p('8 bảng mới; tái sử dụng 10 bảng hiện có. Bảng hiện có chỉ liệt kê cột liên quan. Tất cả bảng có created_at và updated_at kiểu DATETIME; các bảng mới dùng PK BIGINT tự tăng. PK là khóa chính, FK là khóa ngoại, UNIQUE là không trùng.')
for name,state,purpose,fields in schemas:
    # Keep headings readable; long SQL identifiers appear in body at size 14.
    h(name.replace('_',' '))
    p(f'{name} — {state}. Công dụng: {purpose}')
    rows=[]
    # Group adjacent fields with the same data type to reduce redundant table rows.
    for field,typ,note in fields:
        if rows and rows[-1][1]==typ.upper() and len(rows[-1][0])+len(field)<45:
            rows[-1][0]+=', '+field;rows[-1][2]+='; '+note
        else: rows.append([field,typ.upper(),note])
    tab(['Trường','Kiểu dữ liệu','Ý nghĩa và ràng buộc'],rows,[2.6,1.65,2.95])
    if name=='orders':p('UNIQUE(order_group_id, shop_id). customer_id phải khớp group. Tiền không âm; dùng DECIMAL, không dùng FLOAT.')
    if name=='order_items':p('UNIQUE(order_id, product_variant_id). Variant phải thuộc product; quantity nguyên dương.')
    if name=='countries':p('countries hiện tại gắn shop. geo_countries là địa lý chung; nối cấu hình giao hàng bằng country_code.')
    if name=='order_addresses':p('Lưu địa chỉ từng cột và giữ nguyên dữ liệu lúc đặt, kể cả khi Customer thay đổi địa chỉ sau này.')
p('Không xóa dây chuyền lịch sử đơn khi xóa sản phẩm hoặc cấu hình. Tham chiếu dùng RESTRICT và giữ soft delete hiện có. Các trường nullable được ghi rõ; địa chỉ bắt buộc theo quốc gia.')
doc.core_properties.title='Kế hoạch Create Order'
doc.core_properties.author=''
doc.save(ROOT/'ke-hoach-create-order.docx')

nullable={'province_id','ward','house_number','apartment','paid_at','image_url','description','instructions','province_code','province_name'}
indexes={'provinces':['(country_id, code) [unique]'], 'order_groups':['(user_id, request_id) [unique]'], 'orders':['(order_group_id, shop_id) [unique]'], 'order_items':['(order_id, product_variant_id) [unique]']}
refs=[];parts=['// Dán toàn bộ vào dbdiagram.io. Thiết kế đề xuất; bảng hiện có chỉ lấy cột liên quan.\n// customers khác users. countries theo shop; geo_countries là địa lý chung.']
for name,state,purpose,fields in schemas:
    parts.append(f'\nTable {name} {{\n  Note: \'{state}. {purpose}\'')
    for field,typ,note in fields:
        attrs=[]
        if field=='id':attrs=['pk','increment','not null']
        elif name=='shipping_rate_countries':attrs=['pk','not null']
        else:attrs=['null' if field in nullable else 'not null']
        if 'UNIQUE' in note and 'UNIQUE(' not in note:attrs.append('unique')
        dbtype=typ.replace(' unsigned','')
        if name=='countries' and field=='country_code':dbtype='varchar(2)'
        parts.append(f'  {field} {dbtype} [{", ".join(attrs)}, note: \'{note}\']')
        if 'FK ' in note:
            target=note.split('FK ',1)[1].split(';')[0]
            target=target.split(' ')[0]
            cardinal='-' if 'UNIQUE' in note else '>'
            refs.append(f'Ref: {name}.{field} {cardinal} {target}')
    parts+=['  created_at datetime [not null]','  updated_at datetime [not null]']
    if name in indexes:parts+=['  indexes {']+['    '+x for x in indexes[name]]+['  }']
    parts.append('}')
parts+=['\n// Quan hệ khóa ngoại']+refs
(ROOT/'create-order-erd.dbml').write_text('\n'.join(parts)+'\n',encoding='utf-8')
print('DOCX updated; DBML created')
