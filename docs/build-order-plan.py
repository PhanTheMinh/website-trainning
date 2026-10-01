from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
QA = ROOT / 'order-plan-assets'
QA.mkdir(exist_ok=True)
FONT = 'C:/Windows/Fonts/arial.ttf'
BOLD = 'C:/Windows/Fonts/arialbd.ttf'

def diagram(name, size, nodes, edges):
    im = Image.new('RGB', size, 'white')
    d = ImageDraw.Draw(im)
    f = ImageFont.truetype(FONT, 26)
    fb = ImageFont.truetype(BOLD, 28)
    for points, label, xy in edges:
        d.line(points, fill='#63727C', width=3)
        if label:
            box = d.textbbox(xy, label, font=f)
            d.rectangle((box[0]-5,box[1]-3,box[2]+5,box[3]+3), fill='white')
            d.text(xy, label, font=f, fill='#29333A')
    for rect, title, body in nodes:
        d.rounded_rectangle(rect, radius=12, fill='#F4F7F9', outline='#53616B', width=3)
        x1,y1,x2,y2=rect
        d.text((x1+20,y1+15),title,font=fb,fill='black')
        if body:
            d.line([(x1,y1+58),(x2,y1+58)],fill='#D0D7DC',width=2)
            d.multiline_text((x1+20,y1+72),body,font=f,fill='#29333A',spacing=9)
    out=QA / (name+'.png')
    im.save(out)
    return out

flow=diagram('user-api-flow',(1500,1550),[
 ((360,20,1140,120),'Click Checkout',''),
 ((360,185,1140,305),'Lưu checkout và gọi Create Order','checkout_token + version + request_id'),
 ((360,370,1140,480),'Kiểm tra email trong customers',''),
 ((35,545,675,665),'Đã có Customer','Lấy customer_id hiện có'),
 ((825,545,1465,665),'Chưa có Customer','Tạo Customer từ dữ liệu checkout'),
 ((250,740,1250,890),'Đọc dữ liệu hiện tại và tính lại','Variant price × quantity → sub_total\nShipping rate → shipping_fee'),
 ((250,970,1250,1080),'order_total = sub_total + shipping_fee',''),
 ((250,1150,1250,1310),'Transaction tạo đơn','Order + Items + Address + COD\nTrừ kho và hoàn tất checkout'),
 ((250,1380,1250,1500),'Thành công → Order data','GET thông tin đơn đã lưu để hiển thị')
],[
 ([(750,120),(750,185)],'↓',(765,132)),
 ([(750,305),(750,370)],'↓',(765,317)),
 ([(750,480),(750,510),(355,510),(355,545)],'Có',(380,510)),
 ([(750,480),(750,510),(1145,510),(1145,545)],'Không',(1170,510)),
 ([(355,665),(355,705),(750,705),(750,740)],'',(0,0)),
 ([(1145,665),(1145,705),(750,705)],'',(0,0)),
 ([(750,890),(750,970)],'↓',(765,915)),
 ([(750,1080),(750,1150)],'↓',(765,1100)),
 ([(750,1310),(750,1380)],'Commit',(770,1325))
])

core=diagram('erd-orders',(1600,1450),[
 ((20,20,480,210),'users','PK id\nPhiên đăng nhập'),
 ((565,20,1035,210),'CheckOutToken','PK id  |  FK user_id\nitems, version, selections'),
 ((1120,20,1580,260),'customers','PK id  |  UQ email_key\nEmail và thông tin khách\nĐịa chỉ tách từng cột'),
 ((565,365,1035,615),'order_groups','PK id\nFK checkout_id UQ\nFK user_id, customer_id\nrequest_id'),
 ((565,760,1035,1020),'orders','PK id  |  FK group_id\nFK customer_id, shop_id\norder_name, sub_total\nshipping_fee, order_total'),
 ((20,760,480,1050),'order_items','PK id  |  FK order_id\nFK product_id, variant_id\nquantity, unit_price\nline_total, tên và biến thể'),
 ((1120,760,1580,1050),'order_payments','PK id  |  FK order_id UQ\nFK payment_method_id\nmethod_type, amount\nstatus, tên và hướng dẫn'),
 ((565,1170,1035,1420),'order_addresses','PK id  |  FK order_id UQ\nNgười nhận và liên hệ\nQuốc gia, tỉnh, thành phố\nĐường, số nhà, mã bưu chính')
],[
 ([(250,210),(250,300),(690,300),(690,365)],'1 → N',(260,275)),
 ([(800,210),(800,365)],'1 → 0..1',(815,280)),
 ([(1350,260),(1350,440),(1035,440)],'1 → N',(1160,405)),
 ([(800,615),(800,760)],'1 → N',(815,665)),
 ([(1350,440),(1510,440),(1510,690),(950,690),(950,760)],'Customer 1 → N orders',(1080,654)),
 ([(565,870),(480,870)],'1:N',(490,835)),
 ([(1035,870),(1120,870)],'1:1',(1045,835)),
 ([(800,1020),(800,1170)],'1:1',(815,1080))
])

catalog=diagram('erd-catalog-shipping',(1600,1180),[
 ((20,20,480,200),'shops','PK id\nFK owner_user_id'),
 ((565,20,1035,245),'products','PK id  |  FK shop_id\nTên, trạng thái bán'),
 ((1120,20,1580,245),'product_variants','PK id  |  FK product_id\nGiá, tồn kho, SKU'),
 ((20,385,480,610),'shipping_methods','PK id  |  FK shop_id\nTên, code, status'),
 ((565,385,1035,610),'shipping_rates','PK id  |  FK method_id\nfixed_fee\nmin_days, max_days'),
 ((1120,385,1580,610),'shipping_rate_countries','PK và FK rate_id\nPK và FK country_id'),
 ((1120,760,1580,990),'countries hiện có','PK id  |  FK shop_id\nQuốc gia được shop hỗ trợ\ncountry_code → mã ISO'),
 ((20,760,480,990),'geo_countries mới','PK id  |  UQ code\nDanh mục địa lý chung'),
 ((565,760,1035,990),'provinces mới','PK id  |  FK country_id\ncode, name'),
],[
 ([(480,110),(565,110)],'1:N',(490,72)),
 ([(1035,110),(1120,110)],'1:N',(1045,72)),
 ([(250,200),(250,385)],'1:N',(267,272)),
 ([(480,485),(565,485)],'1:N',(490,448)),
 ([(1035,485),(1120,485)],'1:N',(1045,448)),
 ([(1350,610),(1350,760)],'N:1',(1368,673)),
 ([(480,860),(565,860)],'1:N',(490,824))
])

doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.65)
s.left_margin=s.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'),'Arial')
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12.5)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.space_after=Pt(7)
footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Kế hoạch Create Order  |  ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs: r.font.size=Pt(9)

def p(text,style=None): return doc.add_paragraph(text,style)
def h(text): doc.add_heading(text,2)
def page(title): doc.add_page_break();doc.add_heading(title,1)
def bullets(items):
    for text in items: p(text,'List Bullet')
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [7.1/len(headers)]*len(headers)
    for cell,text in zip(t.rows[0].cells,headers):cell.text=text
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');m.append(e)
            pr.append(m);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6EE' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF'));pr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.05
                for run in para.runs:run.font.size=Pt(10.5);run.bold=i==0
    p('')
    return t
def pic(image,width=6.5):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=para.add_run();run.add_picture(str(image),width=Inches(width))
    inline=run._r.xpath('.//wp:docPr')[0];inline.set('descr',image.stem.replace('-',' '))

p('Kế hoạch triển khai Create Order','Title')
p('Checkout Customer Order và Order data','Subtitle')
p('Ngày lập 01 tháng 10 năm 2026')
p('Mục tiêu là thêm button Checkout ở màn checkout. Khi khách bấm, hệ thống dùng bản nháp đã lưu để tìm hoặc tạo Customer theo email, đọc lại giá và cấu hình vận chuyển hiện tại, tạo đơn COD rồi chuyển sang màn Order data. Tài liệu này là kế hoạch triển khai và tiêu chí nghiệm thu trước khi viết chức năng.')
h('1 Phạm vi cho lần triển khai này')
bullets(['Customer là bảng riêng với User. Email checkout được dùng để tra Customer; chưa có thì tạo Customer trong cùng giao dịch tạo đơn.',
'Duyệt danh sách item checkout theo product_variant_id và quantity. Backend đọc giá hiện tại, tính sub_total, shipping_fee và order_total.',
'Giỏ nhiều shop được chia thành một order cho mỗi shop; order_groups gom các order của cùng một lần bấm. Nếu có shop không hợp lệ, chưa tạo đơn nào.',
'Địa chỉ Customer lưu bằng các cột riêng. Địa chỉ giao của order cũng tách cột trong order_addresses để giữ dữ liệu tại lúc đặt.',
'COD tạo trạng thái thanh toán unpaid. Màn Order data hiển thị order_name, địa chỉ giao, vận chuyển, khoảng ngày dự kiến, items và tổng tiền.'])
h('2 Ý nghĩa của dữ liệu realtime')
p('Realtime ở đây là đọc lại dữ liệu trong database tại thời điểm xác nhận đặt hàng, không dùng giá hoặc phí do trình duyệt gửi làm nguồn tính tiền. Quantity và các lựa chọn lấy từ phiên bản checkout khách đã chốt. Không cần WebSocket cho yêu cầu này.')
p('Sau khi order đã tạo, trang chi tiết đọc dữ liệu order đã lưu. Giá, tên hàng, địa chỉ và phương thức của đơn không đổi theo cấu hình shop sau này. Trạng thái đơn và trạng thái thanh toán có thể được cập nhật theo nghiệp vụ.')
h('3 Hiện trạng cần nối vào')
p('Project có Vue và API Express Sequelize; checkout có items, version, shipping_address, shipping_method và payment_method. Model Customer, Order và Province chưa có. Checkout hiện thiếu số nhà và đường; bảng countries hiện tại là cấu hình theo shop. Các phần này cần được bổ sung theo thiết kế bên dưới.')

page('4 Flow từ người dùng đến Order data')
pic(flow,6.3)
p('Nếu thiếu thông tin, hết hàng, phương thức không còn khả dụng hoặc giá thay đổi, API trả lỗi có mã rõ ràng để khách xử lý tại checkout. Chỉ chuyển trang sau khi transaction đã commit.')

page('5 Flow API và cách tính tiền')
for title,body in [
 ('Bước 1 Lưu và xác nhận checkout','Frontend đợi autosave hoàn tất, giữ một request_id cho lần đặt và gửi checkout_token cùng version. Backend lấy User từ session, khóa checkout, kiểm tra sở hữu, phiên bản và yêu cầu đã tạo trước đó.'),
 ('Bước 2 Tìm hoặc tạo Customer','Chuẩn hóa email theo chính sách trim và không phân biệt hoa thường, lưu email_key unique. Tìm Customer bằng email_key; có thì dùng customer_id; chưa có thì tạo từ thông tin checkout. Tạo Customer không tạo tài khoản User và không yêu cầu password.'),
 ('Bước 3 Chuẩn bị dữ liệu Order items','Gộp các dòng trùng variant; kiểm tra variant thuộc đúng product và shop, quantity nguyên dương, sản phẩm và shop còn bán. Đọc giá và tồn kho hiện tại từ product_variants, khóa các variant theo thứ tự id để hạn chế tranh chấp.'),
 ('Bước 4 Tính tiền hàng','Mỗi dòng: line_total = unit_price hiện tại × quantity checkout. Mỗi shop: sub_total = tổng line_total của shop. Đây là dữ liệu chuẩn bị trong bộ nhớ; chưa insert order_items khi chưa có order_id.'),
 ('Bước 5 Tính phí vận chuyển','Lấy shipping_rate_id đã chọn cho từng shop, đọc shipping_rates.fixed_fee và shipping_methods. Kiểm tra rate thuộc shop, method active, quốc gia được hỗ trợ qua shipping_rate_countries. Thu một lần fixed_fee cho mỗi order shop, không nhân theo quantity.'),
 ('Bước 6 Chốt tổng tiền và COD','order_total = sub_total + shipping_fee. Kiểm tra payment_method của từng shop còn active, chưa xóa, đúng owner và type cod. So sánh dữ liệu tính lại với checkout; nếu giá hoặc phí khác thì trả ORDER_REQUOTE_REQUIRED để khách xác nhận lại.'),
 ('Bước 7 Ghi dữ liệu và hoàn tất','Tạo group và orders trước; dùng order_id để insert order_items, order_addresses, order_payments. Trừ tồn kho, đánh dấu checkout hoàn tất, rồi commit. Một bước lỗi thì rollback cả Customer mới, đơn và tồn kho.')]:
    h(title);p(body)

page('6 Ví dụ tính tiền và thời gian giao')
table(['Shop và variant','Giá realtime','Quantity','Thành tiền'],[
 ('A áo đen size M','300.000 VND','2','600.000 VND'),
 ('A quần size L','200.000 VND','1','200.000 VND'),
 ('B giày size 42','1.500.000 VND','1','1.500.000 VND')],[2.8,1.5,1.1,1.7])
table(['Đơn','sub_total','shipping_fee','order_total'],[
 ('Order shop A','800.000','30.000','830.000'),('Order shop B','1.500.000','45.000','1.545.000'),('Cả lần mua','2.300.000','75.000','2.375.000')],[2,1.7,1.7,1.7])
p('Ví dụ dùng VND. Các giá trị là dữ liệu minh họa công thức, không phải giá của sản phẩm trong project. Thuế, voucher, phụ phí COD và chuyển đổi ngoại tệ để giai đoạn sau; hiện chưa cộng các khoản này.')
h('Cách lưu khoảng giao dự kiến')
p('Lấy min_delivery_days và max_delivery_days từ shipping_rates tại lúc tạo. Giai đoạn đầu đề xuất tính theo ngày lịch, mốc là ngày đặt hàng theo Asia/Saigon. Lưu estimated_delivery_from và estimated_delivery_to trên order; ví dụ rate 2 đến 4 ngày thì hiển thị một khoảng ngày. Đây là ước tính, chưa tích hợp lịch làm việc của hãng vận chuyển.')
h('Quy tắc giá và làm tròn')
p('Dùng DECIMAL cho các cột tiền và thư viện decimal hoặc đơn vị tiền nhỏ nhất để tính trên backend. Không cộng trực tiếp số thực JavaScript. VND hiển thị không có phần lẻ; không gửi số đã định dạng có dấu phẩy vào phép tính. Một order chỉ dùng một currency.')
h('Khi giá thay đổi')
p('Ví dụ giá checkout 300.000 nhưng database hiện là 320.000: trả 409 cùng báo giá mới; giao diện cho khách xem và xác nhận. Checkout cập nhật phiên bản mới, rồi gửi lại Create order. Backend vẫn đọc lại lần nữa để bảo đảm giá cuối cùng hợp lệ.')

page('7 Danh mục bảng cần dùng và bổ sung')
table(['Bảng','Tình trạng','Vai trò'],[
 ('users, shops','Hiện có','Phiên đăng nhập, quyền sở hữu checkout và shop bán.'),
 ('products, product_variants','Hiện có','Kiểm tra hàng, đọc giá và tồn kho realtime.'),
 ('CheckOutToken','Hiện có','Nguồn items, quantity, địa chỉ và lựa chọn đã chốt.'),
 ('payment_methods','Hiện có','Cấu hình COD của từng chủ shop.'),
 ('shipping_methods, shipping_rates','Hiện có','Tên vận chuyển, phí và số ngày giao.'),
 ('countries, shipping_rate_countries','Hiện có','Quốc gia giao được của từng shop và bảng nối rate.'),
 ('customers','Tạo mới','Khách hàng riêng biệt với User, tra bằng email_key.'),
 ('geo_countries, provinces','Tạo mới','Danh mục địa lý chung cho địa chỉ Customer và order.'),
 ('order_groups','Tạo mới','Một lần đặt, chống trùng, nối checkout và User.'),
 ('orders, order_items','Tạo mới','Đơn theo shop và các dòng sản phẩm.'),
 ('order_addresses','Tạo mới','Địa chỉ nhận được lưu tách cột tại lúc đặt.'),
 ('order_payments','Tạo mới','COD, số tiền và trạng thái thanh toán từng order.')
],[2.7,1.1,3.3])
p('Tổng cộng 8 bảng mới trong thiết kế này. Chưa cần thêm bảng customer_addresses vì yêu cầu hiện tại là một bộ địa chỉ tách cột trên customers. Nếu sau này cần sổ nhiều địa chỉ, có thể thêm bảng đó mà giữ nguyên lịch sử order_addresses.')
p('Không đổi ý nghĩa countries hiện có trong lần này. geo_countries có mã quốc gia dùng chung, provinces thuộc geo_countries. Nối dữ liệu địa lý và cấu hình vận chuyển bằng country_code đã chuẩn hóa, không dùng country_id của shop làm địa chỉ toàn hệ thống.')

page('8 ERD các bảng giao dịch')
pic(core,7.0)
p('PK là khóa chính; FK là khóa ngoại; UQ là unique. 1:N là một tới nhiều, 1:1 là một tới một, 0..1 là có thể chưa tạo. order_groups.customer_id và orders.customer_id phải cùng Customer; user_id trong group dùng để kiểm tra quyền xem đơn.')
p('Các tham chiếu product, variant, shop, payment và vận chuyển được trình bày ở ERD tiếp theo. Mọi bảng giao dịch lưu thời điểm tạo và cập nhật; không xóa dây chuyền lịch sử đơn khi dữ liệu cấu hình bị xóa.')

page('9 ERD sản phẩm vận chuyển và địa lý')
pic(catalog,7.0)
table(['Quan hệ bổ sung','Số lượng','Ý nghĩa'],[
 ('shops → orders','1:N','Mỗi order thuộc một shop.'),
 ('product_variants → order_items','1:N','Lưu variant_id và dữ liệu tại lúc mua.'),
 ('products → order_items','1:N','Giữ product_id của dòng hàng.'),
 ('payment_methods → order_payments','1:N','Tham chiếu cấu hình và lưu bản chụp COD.'),
 ('shipping_rates → orders','1:N','Mỗi order chọn một rate và lưu phí cùng số ngày.'),
 ('geo_countries → customers và order_addresses','1:N','country_id là địa lý chung.'),
 ('provinces → customers và order_addresses','1:N','province_id tùy quốc gia; kiểm tra thuộc country.')
],[4.1,.6,2.4])

page('10 Trường Customer và địa chỉ tách cột')
p('customers dùng email_key để tìm khách. customer_id là danh tính nghiệp vụ; User là danh tính đăng nhập. Không tự gắn hoặc đổi User chỉ từ email khách nhập. Customer đã có thì dùng lại id; không tự ghi đè hồ sơ bằng địa chỉ của một đơn mới.')
table(['Trường customers','Kiểu đề xuất','Quy tắc'],[
 ('id','BIGINT PK','Tự tăng.'),('email, email_key','VARCHAR(254)','Email hợp lệ; email_key UNIQUE, chuẩn hóa nhất quán.'),
 ('first_name, last_name','VARCHAR(100)','Thông tin Customer.'),('phone','VARCHAR(30)','Chuỗi, giữ mã quốc gia và số 0 đầu.'),
 ('country_id, province_id','BIGINT FK','country bắt buộc; province theo quốc gia.'),
 ('city, ward','VARCHAR(100)','Thành phố và phường/xã, tùy cấu trúc quốc gia.'),
 ('street','VARCHAR(200)','Tên đường hoặc dòng địa chỉ chi tiết.'),('house_number','VARCHAR(50)','Số nhà, có thể chứa chữ.'),
 ('apartment','VARCHAR(100) NULL','Căn hộ/tòa nhà nếu có.'),('postal_code','VARCHAR(20)','Mã bưu chính là chuỗi, theo quốc gia.'),
 ('created_at, updated_at','DATETIME','Thời điểm tạo và cập nhật.')
],[2.8,1.5,2.8])
h('order_addresses giữ địa chỉ tại lúc đặt')
p('Trường: id, order_id UNIQUE FK, recipient_first_name, recipient_last_name, email, phone, country_id, country_code, country_name, province_id nullable, province_code, province_name, city, ward, street, house_number, apartment, postal_code, created_at, updated_at. Các trường text dùng giới hạn tương ứng ở customers. Tên và mã địa lý được lưu để hiển thị lịch sử dù danh mục sau này đổi.')
p('Thông tin nhận của đơn lấy từ checkout, không lấy địa chỉ cũ của Customer để thay thế lựa chọn mới. Địa chỉ tách cột trên các bảng mới; dữ liệu nháp checkout hiện đang lưu dạng TEXT vẫn được ánh xạ bằng normalizer. Bổ sung các trường địa chỉ vào form, API validator và autosave.')

page('11 Trường nhóm đơn và Order')
table(['Trường order_groups','Kiểu','Quy tắc'],[
 ('id','BIGINT PK','Tự tăng.'),('checkout_id','BIGINT FK UNIQUE','Một checkout chỉ tạo một nhóm.'),
 ('user_id, customer_id','BIGINT FK','Người thao tác và Customer được tìm theo email.'),
 ('request_id','UUID','UNIQUE(user_id, request_id).'),('currency','CHAR(3)','VND trong giai đoạn đầu.'),
 ('created_at, updated_at','DATETIME','Thời gian lưu.')
],[2.5,1.5,3.1])
table(['Trường orders','Kiểu','Quy tắc'],[
 ('id','BIGINT PK','Tự tăng.'),('order_group_id, customer_id, shop_id','BIGINT FK','UNIQUE(order_group_id, shop_id).'),
 ('order_code, order_name','VARCHAR(50), VARCHAR(100)','Code UNIQUE; tên ví dụ Order RS000123.'),
 ('status','VARCHAR(20)','Khởi tạo pending.'),('sub_total, shipping_fee, order_total','DECIMAL(18,2)','Không âm; order_total = sub_total + shipping_fee.'),
 ('currency','CHAR(3)','Đồng nhất với group.'),('shipping_method_id, shipping_rate_id','BIGINT FK','Tham chiếu cấu hình đã kiểm tra.'),
 ('shipping_method_name, shipping_method_code','VARCHAR(100), VARCHAR(50)','Lưu tại thời điểm đặt.'),
 ('min_delivery_days, max_delivery_days','INT UNSIGNED','Lưu số ngày từ rate.'),
 ('estimated_delivery_from, estimated_delivery_to','DATE','Khoảng giao dự kiến đã tính.'),
 ('created_at, updated_at','DATETIME','Thời điểm tạo và cập nhật.')
],[2.8,1.5,2.8])
p('Trạng thái order đề xuất: pending, confirmed, shipping, delivered, cancelled. API Create order chỉ khởi tạo pending; thao tác chuyển trạng thái và hoàn kho khi hủy là chức năng riêng cần làm tiếp.')

page('12 Trường Order items Payment và địa lý')
table(['Trường order_items','Kiểu','Ý nghĩa'],[
 ('id, order_id','BIGINT PK và FK','Dòng hàng thuộc order.'),
 ('product_id, product_variant_id','BIGINT FK','Variant phải thuộc product.'),
 ('product_name, sku','VARCHAR(255), VARCHAR(100)','Tên và SKU tại lúc mua.'),
 ('variant_data, image_url','JSON, TEXT NULL','Size/màu và ảnh cho trang Order data.'),
 ('unit_price, line_total','DECIMAL(18,2)','Giá hiện tại và giá × quantity.'),
 ('quantity','INT UNSIGNED','Nguyên dương; UNIQUE(order_id, product_variant_id).')
],[2.8,1.6,2.7])
table(['Trường order_payments','Kiểu','Ý nghĩa'],[
 ('id, order_id','BIGINT PK và FK UNIQUE','Một bản ghi thanh toán cho mỗi order COD.'),
 ('payment_method_id','BIGINT FK','Lấy đúng method của chủ shop.'),
 ('method_type, method_name','VARCHAR(30), VARCHAR(100)','cod và tên hiển thị lúc đặt.'),
 ('description, instructions','TEXT NULL','Hướng dẫn tại thời điểm đặt.'),
 ('amount, currency','DECIMAL(18,2), CHAR(3)','Khớp tổng phải thu của order.'),
 ('status, paid_at','VARCHAR(20), DATETIME NULL','Khởi tạo unpaid; chưa tự chuyển paid.')
],[2.8,1.6,2.7])
p('geo_countries: id PK, code CHAR(2) UNIQUE, name VARCHAR(100), phone_code VARCHAR(8). provinces: id PK, country_id FK, code VARCHAR(20), name VARCHAR(100), UNIQUE(country_id, code). Cả hai có created_at và updated_at; seed từ danh mục địa lý hiện có, chuẩn hóa UK thành GB.')
p('order_items và order_payments cũng có created_at, updated_at. Những cột lịch sử phải được giữ ổn định; cấu hình đang dùng soft delete thì tham chiếu vẫn còn. FK lịch sử dùng RESTRICT hoặc SET NULL theo chính sách xóa, không dùng CASCADE xóa đơn.')

page('13 Hợp đồng API và màn Order data')
h('POST api orders')
p('Yêu cầu xác thực. Body gồm checkout_token, version và request_id. Không nhận customer_id, unit_price, shipping_fee hoặc order_total từ frontend để quyết định kết quả. Giữ request_id khi retry cùng thao tác; cùng key nhưng khác checkout thì trả lỗi xung đột.')
p('{ "checkout_token": "UUID checkout", "version": 5, "request_id": "UUID thao tác đặt" }')
p('Trả 201 khi tạo mới, hoặc 200 cùng kết quả cũ khi gửi lại. Response chứa order_group_id, orders [{ id, order_code, order_name }] và đường dẫn Order data. Kiểm tra idempotency trước lỗi checkout completed để retry thành công không bị chặn.')
h('GET api order groups id và GET api orders id')
p('Đọc đơn đã lưu; kiểm tra order_groups.user_id là User đang đăng nhập. Tra Customer qua email không cấp quyền xem đơn. Trả group totals, customer liên hệ của giao dịch, từng order, address, shipping, khoảng ngày dự kiến, payment và items. Không tính lại giá lịch sử từ product_variants khi GET.')
h('Nội dung Order data')
bullets(['Một shop: hiện order_name và order_code, trạng thái, địa chỉ giao đầy đủ, tên vận chuyển, khoảng ngày giao, COD unpaid và tổng tiền.',
'Nhiều shop: trang /order-groups/:id hiện tất cả đơn con; mỗi đơn có items, shipping fee, payment và thời gian riêng. Link /orders/:id mở một đơn.',
'Mỗi item hiện ảnh, tên, size/màu, SKU nếu cần, quantity, unit_price và line_total. Dữ liệu lấy từ API order đã lưu.',
'Có trạng thái loading, lỗi truy cập/không tìm thấy, thử lại khi lỗi mạng; refresh hoặc mở URL trực tiếp vẫn đọc đúng đơn.'])
table(['Mã lỗi đề xuất','Xử lý giao diện'],[
 ('ADDRESS_INVALID 400','Đánh dấu các trường cần bổ sung.'),('CHECKOUT_CONFLICT 409','Giữ local edits và yêu cầu đồng bộ lại.'),
 ('ITEM_UNAVAILABLE hoặc OUT_OF_STOCK 409','Chỉ rõ variant và số lượng khả dụng.'),
 ('SHIPPING_UNAVAILABLE hoặc PAYMENT_METHOD_UNAVAILABLE 409','Cập nhật lựa chọn; chưa tạo đơn.'),
 ('ORDER_REQUOTE_REQUIRED 409','Hiện báo giá mới để khách xác nhận.'),('401 hoặc 403 hoặc 404','Xử lý phiên đăng nhập và quyền truy cập.')
],[3.7,3.4])

page('14 Các bước triển khai tuần tự')
for title,body in [
 ('1 Chốt tên và ràng buộc dữ liệu','Dùng các bảng và quan hệ trong tài liệu. Ánh xạ tên vật lý theo Sequelize hiện có; tránh từ SQL order bằng tên orders. Chốt email_key, địa chỉ quốc tế và quy tắc nhiều shop.'),
 ('2 Tạo migration và seed địa lý','Tạo geo_countries rồi provinces, customers, order_groups, orders, order_addresses, order_items, order_payments theo thứ tự FK. Tạo unique/index và seed chuẩn hóa. down đảo ngược thứ tự; kiểm tra trên database test trước.'),
 ('3 Tạo model và associations','Tạo các model tương ứng, khai báo belongsTo/hasMany/hasOne trong src/models/index.js. Đối chiếu allowNull, DECIMAL, bảng và timestamps với migration.'),
 ('4 Bổ sung dữ liệu checkout','Thêm street, house_number, apartment, ward vào form và normalizer/validator/autosave. Ánh xạ country/province sang danh mục chung; giữ bản nháp cũ đọc được.'),
 ('5 Viết service tính và tạo đơn','Tách chuẩn bị items, kiểm tra shipping/payment, tìm Customer, so sánh báo giá và tạo transaction. Khóa checkout, variant, rate và method theo thứ tự nhất quán; customer email unique xử lý race bằng rollback và retry đọc.'),
 ('6 Viết validator controller routes','Bổ sung /api/orders và /api/order-groups vào app. Validate body, session, quyền đọc và serialize response. Không trả password hoặc hồ sơ Customer ngoài phạm vi đơn.'),
 ('7 Thêm button Checkout','Nút cuối checkout dùng nhãn Checkout theo yêu cầu. Đợi flush thành công; chặn double click, có trạng thái Creating order, giữ request_id khi lỗi mạng và retry. Khi commit thành công thì điều hướng; chỉ xóa đúng quantity đã mua khỏi giỏ, giữ phần khách thêm sau đó.'),
 ('8 Tạo Order data và kiểm thử','Tạo service gọi GET và route trang group/order. Kiểm thử đơn một shop, nhiều shop, retry, tồn kho, thay đổi giá/phí và theme sáng tối. Chỉ đưa lên môi trường dùng thật khi kiểm thử database và browser đều đạt.')]:
    h(title);p(body)

page('15 Kiểm thử và điều kiện nghiệm thu')
table(['Tình huống','Kết quả cần đạt'],[
 ('Email Customer đã có hoặc chưa có','Dùng đúng Customer; tạo mới chỉ khi chưa có; không ghi đè hồ sơ cũ.'),
 ('Hai request cùng email mới','Một Customer duy nhất nhờ unique; xử lý xung đột có kiểm soát.'),
 ('Một shop và nhiều shop','Đúng số order; đúng subtotal, phí và tổng từng shop.'),
 ('Giá hoặc phí thay đổi','Báo giá lại, chưa tạo đơn với tiền chưa được khách xác nhận.'),
 ('Variant sai product hoặc hết kho','Từ chối; không tạo dòng hàng hoặc trừ kho dở dang.'),
 ('Hai khách mua chiếc cuối cùng','Chỉ một giao dịch thành công; tồn kho không âm.'),
 ('COD hoặc vận chuyển bị tắt/xóa','Từ chối và giữ dữ liệu checkout để khách sửa.'),
 ('Double click và mạng mất sau commit','Một nhóm đơn; retry trả lại kết quả đã tạo.'),
 ('Lỗi giữa transaction','Rollback Customer mới, group, order, items, payment, kho và checkout.'),
 ('Đổi địa chỉ Customer hoặc giá sản phẩm sau đặt','Đơn cũ vẫn hiển thị dữ liệu giao dịch đã lưu.'),
 ('User khác mở URL đơn','Không xem được đơn chỉ nhờ biết email hoặc id.'),
 ('Refresh Order data và hai theme','Đủ thông tin, dễ đọc; xử lý loading, lỗi mạng và retry.')
],[3.1,4.0])
h('Các lựa chọn dùng cho bản triển khai đầu')
p('Đề xuất áp dụng một Customer theo email chuẩn hóa, địa chỉ tách cột; một order mỗi shop; COD unpaid; phí flat theo rate và quốc gia; khoảng giao theo ngày lịch tính từ ngày đặt; tổng tiền chỉ gồm hàng và vận chuyển. Đây là các lựa chọn kỹ thuật của plan để triển khai nhất quán.')
p('Phát triển tiếp: nhiều địa chỉ Customer, thuế/voucher, phụ phí, cổng thanh toán online, vận chuyển theo tỉnh/cân nặng, xác nhận thu COD, hủy đơn và hoàn kho. Chức năng hủy/hoàn kho cần hoàn thiện trước khi vận hành xử lý đơn thực tế.')

doc.core_properties.title='Kế hoạch triển khai Create Order'
doc.core_properties.subject='Checkout Customer Order ERD và kế hoạch triển khai'
doc.core_properties.author=''
out=ROOT/'ke-hoach-create-order.docx'
doc.save(out)
print(out)
print('Sections:', len(doc.sections), 'Tables:', len(doc.tables), 'Diagrams:', len(doc.inline_shapes))
